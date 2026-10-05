"""Regenerate the report with python-docx installed in the project environment."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parent / "Al-Majarrat-Project-Documentation.docx"
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(41, 55, 49)
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size, color in (
    ("Title", 28, "123D32"),
    ("Heading 1", 19, "123D32"),
    ("Heading 2", 14, "26715A"),
    ("Heading 3", 11, "26715A"),
):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(5)
styles["Heading 1"].paragraph_format.page_break_before = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("AL-MAJARRAT  |  PROJECT DOCUMENTATION")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(110, 125, 116)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run(
    "Al-Majarrat | Evidence-Grounded Quran and Hadith Research Assistant  |  Page "
)
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(110, 125, 116)
page_field = OxmlElement("w:fldSimple")
page_field.set(qn("w:instr"), "PAGE")
footer._p.append(page_field)


def paragraph(text="", style=None, italic=False):
    para = doc.add_paragraph(style=style)
    run = para.add_run(text)
    run.italic = italic
    return para


def bullet_list(items):
    for item in items:
        paragraph(item, "List Bullet")


def numbered_list(items):
    for item in items:
        paragraph(item, "List Number")


def code_block(text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.18)
    para.paragraph_format.right_indent = Inches(0.18)
    para.paragraph_format.space_after = Pt(7)
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(31, 70, 56)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F1F5F1")
    para._p.get_or_add_pPr().append(shading)


def table(headers, rows):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.style = "Light Shading Accent 1"
    for index, value in enumerate(headers):
        cell = result.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(value))
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        fill = OxmlElement("w:shd")
        fill.set(qn("w:fill"), "123D32")
        cell._tc.get_or_add_tcPr().append(fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        cells = result.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = ""
            para = cells[index].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.add_run(str(value)).font.size = Pt(8.5)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "EFF5F0")
                cells[index]._tc.get_or_add_tcPr().append(fill)
    doc.add_paragraph()


def note(text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.18)
    run = para.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(92, 109, 99)


def heading(text, level):
    doc.add_heading(text, level=level)


# Cover
for _ in range(3):
    doc.add_paragraph()
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = cover.add_run("AL-MAJARRAT")
run.bold = True
run.font.size = Pt(17)
run.font.color.rgb = RGBColor.from_string("26715A")
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = cover.add_run("PROJECT DOCUMENTATION")
run.bold = True
run.font.size = Pt(29)
run.font.color.rgb = RGBColor.from_string("123D32")
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = cover.add_run(
    "Evidence-Grounded Quran and Hadith\nResearch Assistant"
)
run.font.size = Pt(19)
run.font.color.rgb = RGBColor.from_string("26715A")
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = cover.add_run("A local-first retrieval-augmented generation system")
run.italic = True
run.font.size = Pt(12)
for _ in range(2):
    doc.add_paragraph()
for label in (
    "Prepared by: [Student / Project Team]",
    "Guided by: [Supervisor / Guide]",
    "Institution: [Institution Name]",
    "Department: [Department Name]",
    "Academic year: 2026",
):
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.add_run(label).font.size = Pt(11)
doc.add_page_break()

heading("Abstract", 1)
paragraph(
    "Al-Majarrat is a locally operated research assistant for exploring Quranic "
    "verses and Hadith passages. It combines a structured SQLite corpus, SQLite "
    "FTS5 lexical search, locally generated text embeddings, rank fusion, and an "
    "optional local large language model (LLM). A command-line interface supports "
    "direct search and question answering; a browser interface exposes question "
    "entry, source filters, retrieved references, answer copying, locally retained "
    "conversation history, and a responsive sidebar. The system is designed to "
    "ground generated answers in retrieved source passages, attach reference tags, "
    "and present source text for user verification."
)
paragraph(
    "The implemented data store contains 34,157 Hadith records from six "
    "collections and 56,124 Quran rows representing Arabic text and translation "
    "records. The current vector tables contain 34,157 Hadith embeddings and 6,236 "
    "Quran verse embeddings. These are database coverage counts, not measures of "
    "answer accuracy. The project health test suite contains four tests; all four "
    "passed in the validation run recorded for this report. Data-file and SQLite "
    "schema checks also passed. Ollama-dependent validation was intentionally "
    "skipped during that run, and no independent retrieval-quality benchmark or "
    "user study is claimed."
)
paragraph(
    "This document describes the implemented scope, architecture, data "
    "preparation, retrieval and generation workflow, verified results, limitations, "
    "operation, and future work. Al-Majarrat is a research aid, not a religious "
    "authority; users should verify translations, references, and interpretations "
    "against the cited source editions and qualified scholarship."
)
paragraph(
    "Keywords: Quran, Hadith, SQLite FTS5, semantic search, embeddings, "
    "retrieval-augmented generation, Ollama, local AI."
)
doc.add_page_break()
heading("Contents", 1)
toc = doc.add_paragraph()
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
toc._p.append(field)
note(
    "The table of contents is generated from the document headings. In Word, "
    "right-click it and choose Update Field → Update entire table if you edit "
    "the report."
)

heading("Certificate", 1)
paragraph("[MANUAL INPUT REQUIRED] Certificate of completion / project approval to be inserted here.")
paragraph("This section should contain the official certificate text, student / team name, department, guide name, college / university details, and signatures where applicable.")
paragraph("Placeholders are intentionally retained because the source project does not contain verified certificate data.")
heading("Preface", 1)
paragraph("[MANUAL INPUT REQUIRED] Preface text to be added by the project author or supervisor.")
paragraph("This section may include the purpose of the project, the educational context, and the motivation behind the implementation of Al-Majarrat.")
heading("Acknowledgement", 1)
paragraph("[MANUAL INPUT REQUIRED] Acknowledgement text to be inserted here.")
paragraph("This section should be completed only with names, institutions, supervisors, or mentors that are officially confirmed by the project team.")
heading("Formal Project Submission Details", 1)
table(
    ["Field", "Details"],
    [
        ("Project Title", "Al-Majarrat"),
        ("Submitted by", "[Student / Team Name]"),
        ("Guided by", "[Guide / Supervisor Name]"),
        ("Institution", "[Institution Name]"),
        ("Department", "[Department Name]"),
        ("Submission Type", "[Final / Mid-Year / Semester Project]"),
        ("Academic Year", "[2025-2026 or as applicable]"),
        ("Status", "[MANUAL INPUT REQUIRED]"),
    ],
)
paragraph("Note: Formal submission details are intentionally kept as placeholders because the current repository does not contain verified certificate, guide, or institutional records.")
doc.add_page_break()

# Chapter 1
heading("Chapter 1 — Introduction", 1)
heading("1.1 Background of the Study", 2)
paragraph(
    "The Quran and Hadith are foundational sources for Islamic study. Readers "
    "often need to locate material by a verse reference, topic, wording, collection, "
    "narrator, or translation. Large text collections make manual lookup slow, "
    "while general-purpose language models can produce unsupported statements or "
    "uncertain citations if they answer without a source-grounding step."
)
paragraph(
    "Al-Majarrat addresses this information-retrieval problem by maintaining "
    "local Quran and Hadith corpora and combining lexical and semantic retrieval. "
    "Search results are formatted with references and source text. When the optional "
    "local LLM is available, retrieved passages are supplied as constrained context "
    "for a natural-language response. The system is intended to make exploratory "
    "research more accessible while keeping supporting material visible."
)
heading("1.2 Problem Definition", 2)
paragraph(
    "A user asking a thematic question may not know the exact wording or reference "
    "needed to find relevant source material. Keyword-only search can miss "
    "semantically related phrasing; unconstrained generation can return claims "
    "that are not supported by a verifiable source. There is a need for a searchable, "
    "locally operable system that can retrieve relevant Quran and Hadith passages, "
    "expose their references, and distinguish source evidence from generated "
    "explanation."
)
heading("1.3 Objectives of the Project", 2)
bullet_list(
    [
        "Create a project-root-stable data and source layout for the application.",
        "Prepare Quran and Hadith source files and consolidate them into SQLite.",
        "Provide full-text search using SQLite FTS5 and semantic retrieval using locally generated embeddings.",
        "Combine lexical and semantic rankings to produce relevant passages.",
        "Provide CLI and browser workflows for searching and asking questions.",
        "Constrain local LLM answers to retrieved context and display references for verification.",
        "Provide health checks and tests for paths, source data, and database schema.",
    ]
)
heading("1.4 Scope of the Project", 2)
paragraph(
    "The implemented application covers English-language question handling, "
    "Quran and Hadith retrieval, selected English translations with Uthmani Arabic "
    "text, a local Ollama-compatible generation endpoint, and a local browser UI. "
    "Quran queries include direct verse-key and Surah handling. Hadith data spans "
    "Sahih al-Bukhari, Sahih Muslim, Sunan an-Nasa’i, Sunan Abi Dawud, Jami‘ "
    "at-Tirmidhi, and Sunan Ibn Majah."
)
paragraph(
    "The project does not implement user accounts, a hosted multi-user service, "
    "scholarly adjudication, or a validated fatwa workflow. Citations do not prove "
    "that generated answers are correct. Fine-tuning dataset export and a QLoRA "
    "training script are included as development utilities; the default serving "
    "path uses the configured Ollama chat model, and a trained adapter is not "
    "automatically loaded."
)
heading("1.5 Motivation", 2)
paragraph(
    "The project pairs convenient thematic exploration with visible source "
    "evidence. A local-first approach permits operation without sending each "
    "research query to a hosted LLM provider, subject to the installed local model "
    "and hardware. The aim is assistance in locating and reviewing sources—not "
    "replacing primary texts, established scholarship, or qualified guidance."
)

# Chapter 2
heading("Chapter 2 — Literature Review", 1)
heading("2.1 Existing Approaches / Related Work", 2)
paragraph(
    "Relevant technical approaches include database text search, term-frequency "
    "ranking, dense vector retrieval, and retrieval-augmented generation. They are "
    "complementary: lexical search is effective for exact words and names; embedding "
    "search can surface related concepts with different wording; grounded generation "
    "can summarize retrieved passages conversationally."
)
heading("2.1.1 Full-Text Search", 3)
paragraph(
    "SQLite FTS5 provides an embedded full-text index and relevance ranking. It "
    "suits a local application because it operates alongside the project database "
    "without a separate search server."
)
heading("2.1.2 Dense Semantic Retrieval", 3)
paragraph(
    "An embedding model maps text to numeric vectors. Query/document similarity "
    "can retrieve passages whose wording differs from the query. Al-Majarrat uses "
    "the configured nomic-embed-text Ollama model and cosine similarity over "
    "stored JSON vectors."
)
heading("2.1.3 Retrieval-Augmented Generation", 3)
paragraph(
    "RAG systems retrieve supporting documents and provide them as context to a "
    "language model. Here, a prompt instructs the model to answer from supplied "
    "Quran/Hadith context, retain exact reference tags, include English translation "
    "text, and state when no direct references were found."
)
heading("2.2 Limitations of Existing Approaches", 2)
paragraph(
    "Keyword search depends on token overlap and can be sensitive to spelling and "
    "phrasing. Dense retrieval depends on embedding-model coverage, corpus quality, "
    "and ranking. Generative models may still misread, overgeneralize, or "
    "misattribute supplied passages. Combining methods does not prove relevance "
    "or scholarly soundness."
)
heading("2.3 Research Gap", 2)
paragraph(
    "A practical gap exists between simple reference lookup and conversational "
    "answers traceable to source passages. A local application should support direct "
    "lookup and thematic search, keep source text visible, and avoid depending on a "
    "hosted model. Al-Majarrat implements these features; systematic relevance, "
    "factuality, and usability evaluations remain future work."
)
heading("2.4 Proposed Approach", 2)
paragraph(
    "Prepare the source corpora, build a SQLite database with FTS5 tables, generate "
    "and store embeddings, retrieve lexical and semantic candidate sets, combine "
    "rankings using Reciprocal Rank Fusion (RRF), and optionally submit selected "
    "evidence to a local chat model. CLI and web UI share the RAG module."
)

# Chapter 3
heading("Chapter 3 — System Analysis", 1)
heading("3.1 System Requirements", 2)
heading("3.1.1 Hardware Requirements", 3)
paragraph(
    "The repository declares no formal minimum hardware. Runtime requirements "
    "depend on selected local models; inference can require substantial memory and "
    "may benefit from a compatible GPU. Retrieval-only use avoids the chat model, "
    "but semantic retrieval requires the embedding endpoint unless it falls back "
    "to lexical-only search."
)
heading("3.1.2 Software Requirements", 3)
table(
    ["Component", "Technology"],
    [
        ("Runtime", "Python; verified environment: Python 3.14.7."),
        ("Database", "SQLite with FTS5 virtual tables."),
        ("Embedding service", "Ollama; configured model nomic-embed-text."),
        ("Generation service", "Ollama OpenAI-compatible endpoint; qwen2.5:3b."),
        ("Browser UI", "HTML, CSS, JavaScript; Python standard-library HTTP server."),
        ("Training utilities", "PyTorch, Transformers, PEFT, Accelerate, bitsandbytes."),
        ("Testing", "Python unittest; four project health tests."),
    ],
)
heading("3.1.3 Functional Requirements", 3)
table(
    ["ID", "Requirement", "Status / evidence"],
    [
        ("FR-01", "Search Hadith from CLI and web-backed retrieval.", "Implemented."),
        ("FR-02", "Search Quran text and direct references.", "Implemented."),
        ("FR-03", "Retrieve using lexical and semantic signals.", "FTS5, vectors, cosine, RRF."),
        ("FR-04", "Generate answers through a local LLM.", "Implemented; service required."),
        ("FR-05", "Filter all/Quran/Hadith in web UI.", "Implemented."),
        ("FR-06", "Display retrieved source passages and references.", "Implemented."),
        ("FR-07", "Retain/reopen completed browser conversations.", "localStorage."),
        ("FR-08", "Delete an individual conversation.", "Implemented."),
        ("FR-09", "Report basic data/backend readiness.", "Validator and status API."),
    ],
)
heading("3.1.4 Non-Functional Requirements", 3)
bullet_list(
    [
        "Portability: configured paths resolve from the project root.",
        "Traceability: results and prompt instructions retain references and evidence.",
        "Local operation: web server defaults to loopback; model endpoints are local.",
        "Input bounds: web questions are limited to 2,000 characters; bodies to 16 KB.",
        "Maintainability: code, data, tests, and UI assets are separated.",
        "Reliability: health checks validate required files and database tables.",
    ]
)
heading("3.1.5 User Requirements", 3)
paragraph(
    "The Al-Majarrat system is intended for users who need a local, evidence-based way to search Quran and Hadith material and ask grounded questions. Based on the implemented functionality, the project supports the following user requirements:"
)
bullet_list(
    [
        "Users can run a local search for Quran verses or Hadith text without relying on a remote hosted service.",
        "Users can query either the Quran, Hadith collections, or both together using a single interface.",
        "Users can submit a natural-language question and receive references tied to retrieved source passages.",
        "Users can inspect the underlying source text, translation, reference, and collection metadata before trusting an answer.",
        "Users can work with the project from the CLI or the browser interface without changing the core system design.",
        "Users can view the current database and model readiness through validation or status endpoints.",
        "Users can use the web application to reopen previous conversations and delete individual sessions from local browser storage.",
        "Users can keep the workflow local-first, especially when using the configured Ollama model and local SQLite database.",
    ]
)
heading("3.2 Feasibility Study", 2)
heading("3.2.1 Technical Feasibility", 3)
paragraph(
    "The project uses Python, SQLite/FTS5, local HTTP model APIs, and browser "
    "technologies. A populated database, vector indexes, and passing health tests "
    "demonstrate that principal data components work in the checked environment. "
    "End-to-end generation depends on Ollama and configured models."
)
heading("3.2.2 Economic Feasibility", 3)
paragraph(
    "Local models can avoid per-query hosted LLM fees. Costs include hardware, "
    "storage, electricity, initial downloads, and maintenance. No formal cost "
    "study has been conducted."
)
heading("3.2.3 Operational Feasibility", 3)
paragraph(
    "The CLI supports direct search and interactive questions; the web app provides "
    "a browser workflow. Setup requires corpus/database preparation and Ollama for "
    "full semantic/generative operation. Users should verify AI responses against "
    "displayed sources."
)
heading("3.2.4 Schedule Feasibility", 3)
paragraph(
    "The repository contains extraction, indexing, application, training, and "
    "validation stages. No schedule, effort estimate, or milestone record was "
    "found in the reviewed implementation evidence; no schedule-completion claim "
    "is made."
)
heading("3.3 Methodology Used", 2)
paragraph(
    "The staged workflow is: obtain data, normalize records, build search indexes, "
    "create vectors, retrieve and fuse candidates, compose grounded prompts, and "
    "present results through CLI or web UI. Automated health checks cover project "
    "structure and schema; this report distinguishes these from evaluations not "
    "yet performed."
)

# Chapter 4
heading("Chapter 4 — System Design", 1)
heading("4.1 System Architecture", 2)
paragraph(
    "Al-Majarrat is a single-machine client/server-style application. The Python "
    "web server serves static UI assets and ask/status APIs. The web handler and "
    "CLI call the shared RAG module. SQLite stores source records, full-text "
    "indexes, and serialized vectors. Ollama exposes local embedding and chat APIs."
)
code_block(
    "SOURCE CORPORA (CSV / source APIs)\n"
    "        |\n"
    "        v\n"
    "Corpus extraction and master-database builder\n"
    "        |\n"
    "        v\n"
    "SQLite: hadiths + quran_verses + FTS5 tables\n"
    "        |                                      \\\n"
    "        | FTS lexical candidates                \\ embedding index\n"
    "        +---------- cosine + RRF fusion <--------+\n"
    "                          |\n"
    "                  selected references\n"
    "                    /             \\\n"
    "              CLI client       Web UI /api/ask\n"
    "                    \\             /\n"
    "                 grounded prompt -> local Ollama chat model\n"
    "                          |\n"
    "                answer and source display"
)
paragraph("Figure 4.1 — Logical architecture of implemented components.")
heading("4.2 UML Diagrams", 2)
heading("4.2.1 Use Case Diagram", 3)
code_block(
    "User\n"
    "|\n"
    "  +-- Search Quran / Hadith / all\n"
    "  +-- Submit question in web UI or CLI\n"
    "  +-- Review source references and passages\n"
    "  +-- Copy answer or store conversation\n"
    "  +-- Delete single conversation from local history\n"
    "\n"
    "System\n"
    "  +-- Validate query and source mode\n"
    "  +-- Query SQLite FTS index\n"
    "  +-- Query vector index (if available)\n"
    "  +-- Rank and fuse candidate results\n"
    "  +-- Prompt local Ollama model with evidence\n"
    "  +-- Return answer plus source references\n"
)
heading("4.2.2 Sequence Diagram", 3)
code_block(
    "User -> Web UI: Submit question (mode = all / quran / hadith)\n"
    "Web UI -> serve_web_ui.py: POST /api/ask\n"
    "serve_web_ui.py -> al_majarrat.py: retrieve_hybrid_context(query)\n"
    "al_majarrat.py -> SQLite: FTS search for lexical hits\n"
    "al_majarrat.py -> SQLite: fetch vector embeddings for semantic matches\n"
    "al_majarrat.py -> al_majarrat.py: fuse lexical + semantic rankings\n"
    "al_majarrat.py -> Ollama: build grounded prompt and request answer\n"
    "Ollama --> al_majarrat.py: answer with source-aware response\n"
    "al_majarrat.py --> serve_web_ui.py: quran + hadith + answer\n"
    "serve_web_ui.py --> Web UI: render answer and source cards\n"
    "Web UI --> User: display result and save local history\n"
)
heading("4.2.3 ER Diagram", 3)
code_block(
    "hadiths --------------------< hadiths_fts\n"
    "| id\n"
    "| collection_slug\n"
    "| book_id\n"
    "| chapter_id\n"
    "| hadith_id\n"
    "| narrator\n"
    "| text_english\n"
    "| text_arabic\n"
    "| reference\n"
    "|\n"
    "quran_verses ----------------< quran_fts\n"
    "| id\n"
    "| chapter_number\n"
    "| verse_number\n"
    "| verse_key\n"
    "| text_uthmani\n"
    "| translation_resource_id\n"
    "| translation_text\n"
    "|\n"
    "quran_vector_index (verse_key -> embedding_json)\n"
    "hadith_vector_index (hadith_db_id -> reference + embedding_json)\n"
)
heading("4.2.4 Data Flow Diagram (DFD Level 0)", 3)
code_block(
    "[User / Researcher]\n"
    "        | asks question\n"
    "        v\n"
    "[Al-Majarrat System]\n"
    "        | uses\n"
    "        +--> [SQLite Quran+Hadith Data + FTS + vectors]\n"
    "        +--> [Ollama model endpoint]\n"
    "        v\n"
    "[Answer + references displayed to user]"
)
heading("4.2.5 DFD Level 1", 3)
code_block(
    "Question input -> Query validation -> Hybrid retrieval\n"
    "                         |-- lexical search (FTS5)\n"
    "                         |-- semantic search (vector index)\n"
    "                         +-- ranking and fusion\n"
    "                         -> prompt construction -> local chat model\n"
    "                         -> answer + source metadata -> UI/CLI output\n")
heading("4.2.6 DFD Level 2", 3)
code_block(
    "Query validation\n"
    "  -> sanitize input\n"
    "  -> check mode (all/quran/hadith)\n"
    "  -> enforce request size limits\n"
    "Hybrid retrieval\n"
    "  -> SQLite lexical retrieval\n"
    "  -> vector similarity retrieval\n"
    "  -> RRF fusion\n"
    "  -> top-source selection\n"
    "Grounded answer\n"
    "  -> assemble source evidence\n"
    "  -> prompt with reference-aware constraints\n"
    "  -> local model response\n"
    "  -> render answer and references\n")
heading("4.2.7 Activity Diagram", 3)
code_block(
    "Start -> enter query -> validate query -> choose source mode\n"
    "  -> run lexical + vector search -> rank and fuse -> check results\n"
    "  -> if results found: build evidence prompt -> call Ollama -> show answer\n"
    "  -> else: show no direct matches and source guidance\n"
    "  -> display answer + references -> save to local history -> End")
heading("4.2.8 Collaboration / Communication Diagram", 3)
code_block(
    "User <-> Web UI <-> serve_web_ui.py <-> al_majarrat.py\n"
    "                                         |\n"
    "                                         +-> SQLite database\n"
    "                                         +-> local Ollama embeddings API\n"
    "                                         +-> local Ollama chat API\n")
heading("4.2.9 Component Diagram", 3)
code_block(
    "[web/index.html]\n"
    "       |\n"
    "       v\n"
    "[serve_web_ui.py]\n"
    "       |\n"
    "       +--> [al_majarrat.py: retrieval + RAG logic]\n"
    "       |\n"
    "       +--> [SQLite DB: hadiths, quran_verses, FTS, vectors]\n"
    "       |\n"
    "       +--> [Ollama: embeddings/chat endpoints]\n")
heading("4.3 Database Design", 2)
table(
    ["Table", "Key fields / role"],
    [
        ("hadiths", "id; collection/book/chapter; narrator; Arabic/English text; grading; reference."),
        ("quran_verses", "id; chapter/verse; verse_key; Uthmani text; translation resource/name/author/text."),
        ("hadiths_fts", "FTS5 external-content index over collection, reference, narrator, and texts."),
        ("quran_fts", "FTS5 external-content index over verse key, Arabic and translation text."),
        ("hadith_vector_index", "hadith_db_id unique key, reference, serialized embedding JSON."),
        ("quran_vector_index", "verse_key unique key and embedding JSON for preferred English translation."),
    ],
)
heading("4.4 Modules and Data Flow", 2)
table(
    ["Module", "Responsibility"],
    [
        ("src/extract_*.py", "Acquire/transform Quran and six Hadith collections."),
        ("src/create_master_corpus.py", "Build source tables, indexes, and FTS5 tables."),
        ("src/build_vector_index.py", "Generate embeddings in batches and persist vectors."),
        ("src/al_majarrat.py", "Resolve references, retrieve/fuse sources, prompt, call Ollama, CLI."),
        ("src/search_majarrat.py", "Interactive and command-line lexical search."),
        ("src/serve_web_ui.py + web/index.html", "Serve UI/API, validate requests, report status, render results."),
        ("src/validate_project.py + tests/", "Health checks and project tests."),
        ("export/train scripts", "Create fine-tuning data and optional QLoRA adapter; separate from default serving."),
    ],
)
heading("4.5 Use Cases and Sequence", 2)
heading("4.5.1 Primary Actor and Use Cases", 3)
paragraph(
    "The primary actor is a local researcher. Use cases include searching text or "
    "a reference, selecting Quran/Hadith/all, reviewing sources, copying answers, "
    "reopening saved history, and deleting an individual conversation. The app "
    "does not implement accounts or login."
)
heading("4.5.2 Answer Request Sequence", 3)
paragraph(
    "(1) User submits question and mode. (2) UI sends bounded JSON to /api/ask. "
    "(3) Server validates the request and invokes shared retrieval. (4) Retrieval "
    "obtains an embedding when available, searches FTS/direct references, ranks "
    "vectors, fuses candidates, and returns evidence. (5) Server builds a constrained "
    "prompt and requests local chat completion. (6) API returns answer and sources. "
    "(7) Browser renders and stores the completed conversation locally."
)

# Chapter 5
heading("Chapter 5 — Implementation", 1)
heading("5.1 Technology Stack", 2)
table(
    ["Layer", "Technology / details"],
    [
        ("Application", "Python and standard library HTTP, SQLite, JSON, HTTP client."),
        ("Data store", "SQLite: data/db/al_majarrat_master.db."),
        ("Lexical index", "SQLite FTS5 external-content tables."),
        ("Vector index", "JSON-serialized embeddings; cosine similarity in application code."),
        ("Embedding model", "nomic-embed-text served by Ollama."),
        ("Chat model", "qwen2.5:3b at Ollama OpenAI-compatible endpoint."),
        ("Frontend", "Static HTML/CSS/JavaScript; responsive UI and localStorage."),
        ("Training", "Transformers/PEFT QLoRA utilities; CUDA PyTorch needed."),
    ],
)
heading("5.2 Algorithms and Pseudocode", 2)
heading("5.2.1 Corpus Construction", 3)
code_block(
    "FOR each of six Hadith CSV inputs:\n"
    "  map collection metadata; skip rows without Arabic and English text\n"
    "  append normalized records to master Hadith CSV\n"
    "CREATE SQLite source tables and indexes\n"
    "LOAD Hadith and Quran CSV rows\n"
    "CREATE and populate FTS5 external-content tables"
)
heading("5.2.2 Embedding Indexing", 3)
code_block(
    "SELECT eligible Quran translations and English Hadith texts\n"
    "SKIP records already indexed\n"
    "FOR each batch (up to 64 texts): request local embeddings\n"
    "  verify vector count; store JSON; commit progress\n"
    "REPORT indexed coverage"
)
heading("5.2.3 Hybrid Retrieval", 3)
code_block(
    "query_vector = request_embedding(question)  # may be unavailable\n"
    "resolve direct Quran verse / Surah reference when present\n"
    "collect FTS5 candidates and top cosine-similarity candidates\n"
    "RRF(candidate) = sum(1 / (60 + rank)) over ranked lists\n"
    "return up to 6 Quran passages and 5 Hadith passages\n"
    "format references and English text for the prompt"
)
paragraph(
    "The implementation also includes query normalization and synonym/fuzzy "
    "support. Direct verse queries take precedence for Quran selection. These "
    "retrieval heuristics have not received a full quality evaluation."
)
heading("5.2.4 Grounded Answer Generation", 3)
code_block(
    "system prompt = English + exact citations + no fabrication\n"
    "context = selected passages or explicit no-match statements\n"
    "POST prompt to configured local chat endpoint\n"
    "return answer or explicit offline/error status with retrieved context"
)
heading("5.3 Implementation Steps", 2)
numbered_list(
    [
        "Create or activate a Python virtual environment and install requirements.txt.",
        "Ensure required corpus CSV files exist under data/corpora.",
        "Run the master corpus builder. It deletes and recreates the configured DB; back up data before rebuilding.",
        "Install/start Ollama and pull nomic-embed-text and qwen2.5:3b.",
        "Run the vector index builder; it skips already indexed records and reports coverage.",
        "Run validation and tests.",
        "Launch the CLI or web UI and open the local address printed by the server.",
    ]
)
code_block(
    "python -m pip install -r requirements.txt\n"
    "python src/create_master_corpus.py\n"
    "ollama pull nomic-embed-text\n"
    "ollama pull qwen2.5:3b\n"
    "python src/build_vector_index.py\n"
    "python src/validate_project.py --skip-ollama\n"
    "python src/serve_web_ui.py"
)
heading("5.4 Project Folder Structure", 2)
code_block(
    "Project AL Majarrat/\n"
    "  project_config.py       shared paths and model defaults\n"
    "  src/                    application, extraction, indexing, validation\n"
    "  data/corpora/           source/master CSV files\n"
    "  data/db/                SQLite database\n"
    "  data/datasets/          fine-tuning JSONL outputs\n"
    "  data/raw/               raw data area\n"
    "  data/models/            model artifacts area\n"
    "  tests/                  project health tests\n"
    "  web/                    browser interface\n"
    "  docs/                   project documentation"
)
heading("5.5 Screenshots / Operational User Interface", 2)
paragraph(
    "CLI examples: python src/al_majarrat.py --ask \"What does the Quran say "
    "about patience?\"; python src/al_majarrat.py --test-retrieval \"2:255\"; "
    "python src/search_majarrat.py. The web app defaults to 127.0.0.1:8765 and "
    "supports all/Quran/Hadith filters, suggested queries, answer copying, saved "
    "conversations, and per-conversation deletion. History is browser-local and "
    "is not synchronized to another device."
)
paragraph("Figure 5.1 — Browser interface in its responsive mobile layout.")
image_paragraph = doc.add_paragraph()
image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
image_paragraph.add_run().add_picture(
    str(OUTPUT.parent / "al-majarrat-web-ui.png"),
    width=Inches(3.1),
)
heading("5.6 Source Code / Important Implementation", 2)
table(
    ["File", "Purpose / role", "Inputs", "Primary output"],
    [
        ("project_config.py", "Stores project paths, DB path, model endpoints, and translation preference.", "Project root and environment flags.", "Stable absolute paths for all modules."),
        ("src/create_master_corpus.py", "Builds the master Hadith CSV and SQLite database, including FTS5 tables.", "CSV corpus files.", "Database with hadiths, quran_verses, and search indexes."),
        ("src/build_vector_index.py", "Generates embedding batches and persists vector JSON for Quran and Hadith texts.", "Eligible text and Ollama embedding endpoint.", "quran_vector_index and hadith_vector_index tables."),
        ("src/al_majarrat.py", "Core retrieval and generation logic; resolves direct references, searches, ranks, and prompts the local model.", "User query and database state.", "Retrieved evidence and grounded answer."),
        ("src/search_majarrat.py", "Interactive CLI search for Quran and Hadith records using SQLite search functions.", "User query or command input.", "Matched source passages and references."),
        ("src/serve_web_ui.py", "Hosts the browser UI and REST endpoints, including /api/status and /api/ask.", "HTTP requests and DB/model status.", "JSON answer payload plus status reporting."),
        ("web/index.html", "Browser interface for prompt input, source display, conversation history, and copy actions.", "User interaction and API responses.", "Rendered UI with dynamic answer card."),
        ("src/validate_project.py", "Checks corpus completeness, database readiness, and Ollama connectivity.", "Project files and local services.", "Validation status and error messages."),
        ("tests/test_project_health.py", "Regression tests for data presence, DB validity, and missing-file behavior.", "Python test runner.", "Pass/fail result for project health."),
    ],
)
paragraph(
    "The most important implementation logic remains in the shared retrieval module and the UI server. These components are the actual bridge between project data, evidence selection, and answer generation."
)

# Chapter 6
heading("Chapter 6 — Results and Discussion", 1)
heading("6.1 Sample Input and Output", 2)
paragraph(
    "The application supports real query execution through both the CLI and the browser interface. A representative request follows the pattern below, using the implemented system rather than a fabricated benchmark:"
)
code_block(
    "Sample query: What does the Quran say about patience?\n"
    "Command: python src/al_majarrat.py --ask \"What does the Quran say about patience?\"\n"
    "Processing path: query validation -> lexical + semantic retrieval -> source ranking -> prompt assembly -> local model response (if available)\n"
    "Output shape: {\n"
    "  \"query\": \"...\",\n"
    "  \"mode\": \"all\",\n"
    "  \"quran\": [verses and references],\n"
    "  \"hadith\": [hadith references],\n"
    "  \"answer\": \"...\"\n"
    "}\n"
    "Important note: the system displays retrieved source references and evidence; it does not guarantee a scholar-verified answer without human review."
)
heading("6.2 Graphs and Charts", 2)
paragraph(
    "No graphs or charts have been generated for this report because the project does not include a validated retrieval benchmark, user survey, or standardized evaluation dataset. The available project evidence is numerical and operational rather than comparative performance data."
)
paragraph("[REAL DATA REQUIRED] If a future evaluation is performed, charts may be added for corpus size, retrieval coverage, query latency, or citation precision.")
heading("6.3 Analysis of Results", 2)
paragraph(
    "Counts below were read from the populated SQLite database during report "
    "preparation. They describe stored rows and vector coverage, not answer accuracy."
)
table(
    ["Measure", "Count", "Interpretation"],
    [
        ("Hadith source rows", "34,157", "Six loaded collections."),
        ("Quran database rows", "56,124", "Arabic and translation rows; not unique verses."),
        ("Distinct Quran verse keys", "6,236", "Unique verse references in this database."),
        ("Hadith vector rows", "34,157", "Count matches Hadith source row total."),
        ("Quran vector rows", "6,236", "Count matches unique Quran verse keys."),
        ("FTS Hadith rows", "34,157", "External-content index rows."),
        ("FTS Quran rows", "56,124", "External-content index rows."),
    ],
)
table(
    ["Hadith collection", "Rows"],
    [
        ("Sahih al-Bukhari", "7,580"),
        ("Sahih Muslim", "7,360"),
        ("Sunan an-Nasa’i", "5,679"),
        ("Sunan Abi Dawud", "5,272"),
        ("Jami‘ at-Tirmidhi", "3,926"),
        ("Sunan Ibn Majah", "4,340"),
        ("Total", "34,157"),
    ],
)
paragraph(
    "The current database has vector coverage for all Hadith rows and for all distinct Quran verse keys in the snapshot. Health tests confirm the presence of required files, tables, and database structure. These results support operational readiness of the local data layer but do not validate answer quality or retrieval accuracy."
)
paragraph(
    "The live Ollama chat endpoint was intentionally skipped in the recorded health run, so the project cannot claim verified live-model inference without an explicit environment check. Similarly, no retrieval-quality benchmark or user study was performed. The system therefore remains a functional local research assistant with evidence-aware retrieval, but not a formally benchmarked academic evaluation system."
)

# Chapter 7
heading("Chapter 7 — Conclusion and Future Scope", 1)
heading("7.1 Summary of Work", 2)
paragraph(
    "Al-Majarrat provides a local research workflow over Quran and Hadith corpora. "
    "It combines corpus processing, populated SQLite/FTS5 tables, vector indexes, "
    "direct reference lookup, hybrid retrieval with RRF, prompt-constrained local "
    "generation, CLI and web interfaces, health validation, browser history, and a "
    "responsive UI. Code, data, tests, web assets, and docs are separated."
)
heading("7.2 Achievements", 2)
bullet_list(
    [
        "Prepared a database with 34,157 Hadith records and 56,124 Quran rows.",
        "Populated vectors for 34,157 Hadiths and 6,236 distinct Quran verse keys.",
        "Implemented lexical and semantic retrieval with reciprocal-rank fusion.",
        "Added direct Quran reference and Surah handling.",
        "Implemented a local web UI/API and CLI.",
        "Added browser conversation persistence, reopen, and individual deletion.",
        "Passed all four project health tests and configured corpus/database checks.",
    ]
)
heading("7.3 Limitations", 2)
numbered_list(
    [
        "Generated answers may be wrong or overstate evidence; prompt rules cannot guarantee factuality.",
        "No expert-labeled retrieval benchmark, factuality evaluation, performance study, or user survey was conducted.",
        "The live Ollama chat endpoint was not checked in the recorded health run.",
        "Translation, edition, grading, provenance, and license details should be verified before scholarly use or redistribution.",
        "Browser history lives in one browser's localStorage; clearing browser data removes it.",
        "The RAG workflow currently targets English answers; Arabic-first interaction is not its documented focus.",
        "Vector storage/query scaling and latency have not been profiled.",
        "The training utilities do not automatically load their adapter into default inference.",
    ]
)
heading("7.4 Future Enhancements", 2)
bullet_list(
    [
        "Build an expert-reviewed evaluation set; report Recall@k, MRR, citation precision, and unsupported claims.",
        "Record source provenance, edition/version metadata, license notices, and refresh procedures.",
        "Evaluate retrieval filters, rank fusion, reranking, and translation-specific indexing.",
        "Add a model readiness/smoke-test workflow and reproducible runtime benchmarks.",
        "Add history export/import, retention controls, and privacy documentation.",
        "Improve deployment hardening, logging, error reporting, and operational guidance before non-local exposure.",
        "Profile response latency and memory on defined hardware.",
        "If fine-tuning is pursued, add adapter loading and held-out evaluation.",
    ]
)

# References and appendices
heading("Chapter 8 — References", 1)
paragraph("Technical references reviewed for this report (accessed 4 October 2026):")
for reference in (
    '[1] SQLite, “FTS5 Extension.” https://www.sqlite.org/fts5.html',
    '[2] Ollama, “OpenAI compatibility.” https://docs.ollama.com/api/openai-compatibility',
    '[3] Ollama, “Embeddings.” https://docs.ollama.com/capabilities/embeddings',
    "[4] Quran Foundation, Content API documentation. https://api-docs.quran.foundation/",
    "[5] fawazahmed0, “Hadith API” dataset repository. https://github.com/fawazahmed0/hadith-api",
    "[6] Qwen, Qwen2.5-3B-Instruct model card. https://huggingface.co/Qwen/Qwen2.5-3B-Instruct",
    '[7] Python Software Foundation, “sqlite3.” https://docs.python.org/3/library/sqlite3.html',
    "[8] Project files reviewed: project_config.py; src/al_majarrat.py; corpus, index, search, web, validation and training scripts; tests/test_project_health.py.",
    "[9] Al-Majarrat README, setup and operational instructions.",
):
    paragraph(reference)
note(
    "Dataset/API terms and model licenses may vary by source or change over time. "
    "Verify current terms, attribution, and translation permissions before "
    "publishing or redistributing corpus data."
)

heading("Chapter 9 — Appendix", 1)
heading("9.1 Database Structure Summary", 2)
code_block(
    "hadiths(id, collection_slug, collection_name_en/ar, book_id/name, chapter_id/name,\n"
    "        hadith_id, hadith_id_in_book, narrator, text_english, text_arabic, grading, reference)\n"
    "quran_verses(id, chapter_number, verse_number, verse_key, text_uthmani,\n"
    "             translation_resource_id, translation_name, translation_author, translation_text)\n"
    "hadiths_fts: FTS5 external-content index -> hadiths.id\n"
    "quran_fts: FTS5 external-content index -> quran_verses.id\n"
    "hadith_vector_index(id, hadith_db_id UNIQUE, reference, embedding_json)\n"
    "quran_vector_index(id, verse_key UNIQUE, embedding_json)"
)
heading("9.2 Local API Summary", 2)
table(
    ["Method / route", "Purpose"],
    [
        ("GET /", "Serve web/index.html."),
        ("GET /api/status", "Return database counts and best-effort Ollama availability."),
        ("POST /api/ask", 'Accept {"query":"...", "mode":"all|quran|hadith"} and return answer/sources.'),
    ],
)
paragraph(
    "POST /api/ask enforces a non-empty query, a 2,000-character query limit, "
    "allowed source modes, and a 16 KB request body limit. The server defaults "
    "to 127.0.0.1."
)
heading("9.3 Validation and Test Evidence", 2)
code_block(
    "python src/validate_project.py --skip-ollama\n"
    "python -m unittest tests/test_project_health.py"
)
paragraph(
    "Recorded result: source checks and database validation passed; the validator "
    "reported 34,157 Hadith rows and 56,124 Quran rows; all four unittest cases "
    "passed. Ollama was skipped, so this does not certify a live model request."
)
heading("9.4 Main Project Files", 2)
table(
    ["Path", "Description"],
    [
        ("project_config.py", "Shared paths, DB path, model endpoints, preferred translation."),
        ("src/al_majarrat.py", "Hybrid RAG, prompt composition, local generation, CLI."),
        ("src/create_master_corpus.py", "Master corpus and SQLite/FTS builder."),
        ("src/build_vector_index.py", "Embedding generation and persistence."),
        ("src/search_majarrat.py", "Lexical Quran/Hadith CLI search."),
        ("src/serve_web_ui.py", "Local HTTP server, status endpoint, ask API."),
        ("web/index.html", "Responsive UI and browser-local conversation history."),
        ("src/validate_project.py and tests/", "Health checks and unit tests."),
        ("src/export_finetuning_dataset.py; src/train_qwen_lora.py", "Dataset export and optional QLoRA utilities."),
    ],
)
heading("9.5 User Operation Quick Reference", 2)
code_block(
    "Start web UI:       python src/serve_web_ui.py\n"
    "Ask a question:     python src/al_majarrat.py --ask \"What does the Quran say about patience?\"\n"
    "Test retrieval:     python src/al_majarrat.py --test-retrieval \"2:255\"\n"
    "Lexical search:     python src/search_majarrat.py\n"
    "Validate data/DB:   python src/validate_project.py --skip-ollama\n"
    "Full model check:   python src/validate_project.py"
)

update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
doc.settings.element.append(update_fields)
doc.core_properties.title = "Al-Majarrat Project Documentation"
doc.core_properties.subject = (
    "Technical documentation for the Al-Majarrat Quran and Hadith research assistant"
)
doc.core_properties.author = "[Student / Project Team]"
doc.core_properties.keywords = "Al-Majarrat, Quran, Hadith, RAG, SQLite FTS5, Ollama"
doc.save(OUTPUT)
print(f"Created {OUTPUT} ({OUTPUT.stat().st_size:,} bytes)")
